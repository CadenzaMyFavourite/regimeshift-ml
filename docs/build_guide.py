# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx>=1.1,<2"]
# ///
"""Build the Word companion from PROJECT_GUIDE.md, without scientific dependencies.

Optional maintenance command: uv run --no-project docs/build_guide.py
The Markdown remains the editable source of truth. Inspect a rendered copy after edits.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "PROJECT_GUIDE.md"
DESTINATION = HERE / "RegimeShift_ML_Project_Guide.docx"
REPOSITORY = "https://github.com/CadenzaMyFavourite/regimeshift-ml/blob/main/docs/"


def element(tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn(f"w:{key}"), str(value))
    return node


def inline(paragraph, content):
    pattern = r"(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
    for part in re.split(pattern, content):
        if not part:
            continue
        if part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.2)
        elif part.startswith("[") and "](" in part:
            label, address = re.match(r"\[(.*?)\]\((.*?)\)", part).groups()
            if not address.startswith("https://"):
                address = REPOSITORY + address
            relation = paragraph.part.relate_to(
                address,
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                is_external=True,
            )
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), relation)
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            properties.append(element("w:color", val="245571"))
            run.append(properties)
            text = OxmlElement("w:t")
            text.text = label
            run.append(text)
            link.append(run)
            paragraph._p.append(link)
        else:
            paragraph.add_run(part)


def add_table(document, rows):
    count = len(rows[0])
    widths = [2.6, 4.3] if count == 2 else [1.65, 2.25, 3.0]
    if rows[0][0] == "Paired comparison":
        widths = [2.6, 1.3, 3.0]
    if rows[0][0].startswith("Block"):
        widths = [2.5, 1.0, 3.4]
    table = document.add_table(rows=1, cols=count)
    table.autofit = False
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(element(f"w:{side}", val="single", sz="4", color="D9D9D9"))
    props.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        margins.append(element(f"w:{side}", w="95", type="dxa"))
    props.append(margins)
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for row_index, cells in enumerate(rows):
        row = table.rows[0] if row_index == 0 else table.add_row()
        row._tr.get_or_add_trPr().append(element("w:cantSplit"))
        if row_index == 0:
            row._tr.get_or_add_trPr().append(element("w:tblHeader"))
        for i, text in enumerate(cells):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                cell._tc.get_or_add_tcPr().append(element("w:shd", fill="DFEAF1"))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            inline(p, text)
            for run in p.runs:
                run.font.size = Pt(9.5)
                run.bold = row_index == 0
            if re.fullmatch(r"[+\-\d.,% ]+", text):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
    section.header_distance = section.footer_distance = Inches(0.3)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(7)
    for name, size in [("Title", 26), ("Subtitle", 12), ("Heading 1", 18),
                       ("Heading 2", 13), ("Heading 3", 11.5), ("Caption", 9)]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(15 if "Heading" in name else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
        for color in style.element.xpath(".//w:color"):
            for attr in ("themeColor", "themeTint", "themeShade"):
                color.attrib.pop(qn(f"w:{attr}"), None)
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    header = section.header.paragraphs[0]
    header.add_run("REGIMESHIFT ML  |  RESEARCH AND ENGINEERING GUIDE").font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Version 0.1.0  •  ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "RegimeShift ML Research and Engineering Guide"
    doc.core_properties.subject = "Purpose, methods, architecture, and implementation evidence"
    doc.core_properties.author = "RegimeShift ML project"
    doc.core_properties.keywords = "synthetic validation, causal replay, adaptive model weights"
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            for j, code_line in enumerate(code):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0 if j < len(code) - 1 else 8)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.left_indent = Inches(0.12)
                p.paragraph_format.keep_with_next = j < len(code) - 1
                run = p.add_run(code_line or " ")
                run.font.name = "Consolas"
                run.font.size = Pt(8.7)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.fullmatch(r"[| :\-]+", lines[i]):
                    rows.append([x.strip() for x in lines[i].strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue
        elif line.startswith("!["):
            label, path = re.match(r"!\[(.*?)\]\((.*?)\)", line).groups()
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            picture = p.add_run().add_picture(str(HERE / path), width=Inches(6.65))
            picture._inline.docPr.set("descr", label)
            caption = doc.add_paragraph(label + ". Both 95% intervals include zero.")
            caption.style = "Caption"
        elif line.startswith("# "):
            doc.add_paragraph(line[2:], "Title")
        elif line.startswith("## "):
            p = doc.add_paragraph(line[3:], "Heading 1")
            if line.startswith("## 13 "):
                p.paragraph_format.page_break_before = True
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], "Heading 2")
        elif line.startswith("- "):
            inline(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            style = "Subtitle" if i in (2, 4) else "Normal"
            p = doc.add_paragraph(style=style)
            inline(p, line)
            if line.endswith(":"):
                p.paragraph_format.keep_with_next = True
        i += 1
    doc.save(DESTINATION)
    print(f"Created {DESTINATION.name} from {len(SOURCE.read_text().split())} source words")


if __name__ == "__main__":
    build()
